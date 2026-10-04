from __future__ import annotations

import io
import re
from html import unescape
from urllib.parse import urljoin, urlparse, urlunparse

from bs4 import BeautifulSoup

from common.http import HttpClient
from common.models import BaseScraper, JobPayload
from common.normalize import guess_contract, guess_profession, html_to_text

BASE = "https://www.gov.br/inca/pt-br"
CONCURSO_INDEX = f"{BASE}/acesso-a-informacao/institucional/concurso-publico"
CONCURSO_YEARS = (
    f"{CONCURSO_INDEX}/2025",
)
ENSINO_PAGES = (
    f"{BASE}/assuntos/ensino/residencias/medica",
    f"{BASE}/assuntos/ensino",
)
NEWS_SEARCH = f"{BASE}/@@search"
# Página MGI com XLSX de cargos/salários do CPNU 2 (descoberta estável).
CPNU2_CARGOS_PAGE = (
    "https://www.gov.br/gestao/pt-br/concursonacional/cpnu-2/cargos-e-salarios-cpnu-2"
)

OPEN_HINT_RE = re.compile(
    r"edital|sele[cç][aã]o|concurso|inscri[cç]|vaga|resid[eê]ncia|fellow|aperfei[cç]",
    re.I,
)
SKIP_RE = re.compile(r"cookie|privacidade|acessibilidade|mapa\s+do\s+site", re.I)
CPNU_RE = re.compile(
    r"concurso\s+p[uú]blico\s+nacional\s+unificado|\bcpnu\b",
    re.I,
)
INCA_ORG_RE = re.compile(
    r"instituto\s+nacional\s+de\s+c[aâ]ncer|\binca\b",
    re.I,
)
ATTACHMENT_EXT_RE = re.compile(r"\.(pdf|xlsx?|csv|ods)(?:$|[?#])", re.I)
LOADING_TAB_RE = re.compile(r"aguarde\.?\s*carregando", re.I)
MONEY_RE = re.compile(r"R\$\s*[\d.]+(?:,\d{2})?")


def _strip_view_suffix(url: str) -> str:
    """Plone /view em anexos → URL directa do ficheiro."""
    parsed = urlparse(url)
    path = parsed.path.rstrip("/")
    if path.endswith("/view"):
        path = path[: -len("/view")]
    return urlunparse(parsed._replace(path=path))


def _fmt_brl(value: float | int) -> str:
    n = float(value)
    text = f"{n:,.2f}"
    return "R$ " + text.replace(",", "X").replace(".", ",").replace("X", ".")


def parse_cpnu_cargos_xlsx(
    data: bytes,
    *,
    org_re: re.Pattern[str] = INCA_ORG_RE,
) -> list[dict]:
    """Lê cargos/salários do XLSX MGI (coluna Órgão com forward-fill)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    for i, row in enumerate(rows[:12]):
        cells = [str(c or "").strip().lower() for c in row]
        if any("órgão" in c or "orgao" in c for c in cells) and any(
            "cargo" in c for c in cells
        ):
            header_idx = i
            break
    if header_idx is None:
        return []

    header = [str(c or "").strip().lower() for c in rows[header_idx]]

    def col(*needles: str) -> int | None:
        for i, h in enumerate(header):
            if any(n in h for n in needles):
                return i
        return None

    i_org = col("órgão", "orgao")
    i_cargo = col("cargo")
    i_nivel = col("nível", "nivel")
    i_vagas = col("quantidade", "vagas")
    i_rem80 = col("80")
    i_rem100 = col("100")
    if i_org is None or i_cargo is None:
        return []

    out: list[dict] = []
    current_org = ""
    for row in rows[header_idx + 1 :]:
        if not row:
            continue
        org_raw = row[i_org] if i_org < len(row) else None
        if org_raw and str(org_raw).strip():
            current_org = str(org_raw).strip()
        if not current_org or not org_re.search(current_org):
            continue
        cargo = row[i_cargo] if i_cargo < len(row) else None
        if not cargo or not str(cargo).strip():
            continue
        vagas = None
        if i_vagas is not None and i_vagas < len(row):
            try:
                vagas = int(float(row[i_vagas]))  # type: ignore[arg-type]
            except (TypeError, ValueError):
                vagas = None
        rem80 = rem100 = None
        if i_rem80 is not None and i_rem80 < len(row):
            try:
                rem80 = float(row[i_rem80])  # type: ignore[arg-type]
            except (TypeError, ValueError):
                rem80 = None
        if i_rem100 is not None and i_rem100 < len(row):
            try:
                rem100 = float(row[i_rem100])  # type: ignore[arg-type]
            except (TypeError, ValueError):
                rem100 = None
        nivel = None
        if i_nivel is not None and i_nivel < len(row) and row[i_nivel]:
            nivel = str(row[i_nivel]).strip()
        out.append(
            {
                "org": current_org,
                "cargo": str(cargo).strip(),
                "nivel": nivel,
                "vagas": vagas,
                "rem80": rem80,
                "rem100": rem100,
            }
        )
    return out


def format_cargo_summary(cargos: list[dict]) -> str:
    if not cargos:
        return ""
    total = sum(c["vagas"] or 0 for c in cargos)
    lines = [
        "Cargos e remunerações (quadro oficial CPNU 2 — valores sem auxílios):",
        f"Total de vagas INCA: {total}." if total else "Vagas INCA no quadro oficial:",
    ]
    for c in cargos:
        bits = [c["cargo"]]
        if c.get("nivel"):
            bits.append(c["nivel"])
        if c.get("vagas") is not None:
            bits.append(f"{c['vagas']} vaga(s)")
        rem_parts = []
        if c.get("rem80") is not None:
            rem_parts.append(f"{_fmt_brl(c['rem80'])} (80 pts GD)")
        if c.get("rem100") is not None:
            rem_parts.append(f"{_fmt_brl(c['rem100'])} (100 pts GD)")
        if rem_parts:
            bits.append(" / ".join(rem_parts))
        lines.append("- " + " — ".join(bits))
    return "\n".join(lines)


def salary_from_cargos(cargos: list[dict]) -> str | None:
    vals: list[float] = []
    for c in cargos:
        for key in ("rem80", "rem100"):
            v = c.get(key)
            if isinstance(v, (int, float)):
                vals.append(float(v))
    if not vals:
        return None
    lo, hi = min(vals), max(vals)
    if abs(lo - hi) < 0.01:
        return _fmt_brl(lo)
    return f"{_fmt_brl(lo)} a {_fmt_brl(hi)}"


def title_from_cargos(base_title: str | None, cargos: list[dict]) -> str:
    """Título SEO: cargos principais + CPNU/INCA."""
    names: list[str] = []
    seen: set[str] = set()
    for c in cargos:
        raw = c["cargo"]
        short = re.split(r"\s*\(", raw, maxsplit=1)[0].strip()
        # Normaliza variantes Pesquisador (RT…)
        key = re.sub(r"\s+", " ", short).casefold()
        if key in seen:
            continue
        seen.add(key)
        names.append(short)
    if not names:
        return (base_title or "Concurso INCA").strip()
    if len(names) == 1:
        role = names[0]
    elif len(names) == 2:
        role = f"{names[0]} e {names[1]}"
    else:
        # Mantém ≤70 chars para não passar pelo titleCase do shortenJobTitle.
        role = f"{names[0]} e outras especialidades"
    if base_title and CPNU_RE.search(base_title):
        composed = f"{role} — INCA / CPNU 2"
    else:
        composed = f"{role} — INCA"
    return composed[:200]


class IncaScraper(BaseScraper):
    """INCA — analogia IPO (oncologia pública). Editais event-driven no gov.br."""

    slug = "inca"
    name = "INCA — Instituto Nacional de Câncer"
    enrich_details: bool = True
    max_detail_fetches: int = 25

    def fetch(self) -> list[JobPayload]:
        client = HttpClient(min_interval=0.45)
        try:
            candidates = self._collect_candidates(client)
            details: dict[str, dict] = {}
            if self.enrich_details:
                for item in candidates[: self.max_detail_fetches]:
                    try:
                        details[item["url"]] = self._fetch_detail(client, item["url"])
                    except Exception as exc:  # noqa: BLE001
                        print(f"[{self.slug}] detalhe {item['url']}: {exc}")

            jobs: list[JobPayload] = []
            for item in candidates:
                detail = details.get(item["url"]) or {}
                title = (detail.get("title") or item["title"]).strip()
                if not title or SKIP_RE.search(title):
                    continue
                description = detail.get("description") or item.get("summary") or title
                if not OPEN_HINT_RE.search(f"{title} {description[:400]}"):
                    continue
                contract = guess_contract(f"{title} {description}") or "Concurso"
                source_id = item["url"].rstrip("/").split("/")[-1][:160]
                profession_blob = f"{title} {description[:1200]}"
                specialty = detail.get("specialty")
                if not specialty and re.search(r"oncol|c[aâ]ncer|\binca\b", title, re.I):
                    specialty = "Oncologia"
                jobs.append(
                    JobPayload(
                        title=title[:200],
                        company="INCA — Instituto Nacional de Câncer",
                        location_district="Rio de Janeiro",
                        location_concelho="Rio de Janeiro",
                        profession=guess_profession(profession_blob),
                        specialty=specialty,
                        sector="publico",
                        contract_type=contract,
                        description=description[:8000],
                        requirements=detail.get("requirements"),
                        salary=detail.get("salary"),
                        application_url=item["url"],
                        source=self.slug,
                        source_id=source_id,
                        published_at=detail.get("published_at"),
                        expires_at=None,
                    )
                )
            return jobs
        finally:
            client.close()

    def _collect_candidates(self, client: HttpClient) -> list[dict]:
        by_url: dict[str, dict] = {}

        for url in (CONCURSO_INDEX, *CONCURSO_YEARS):
            try:
                html = client.get_text(url)
            except Exception as exc:  # noqa: BLE001
                print(f"[{self.slug}] index {url}: {exc}")
                continue
            for item in self._parse_cards(html, url):
                by_url[item["url"]] = item

        for url in ENSINO_PAGES:
            try:
                html = client.get_text(url)
            except Exception as exc:  # noqa: BLE001
                print(f"[{self.slug}] ensino {url}: {exc}")
                continue
            for item in self._parse_content_links(html, url):
                by_url[item["url"]] = item

        try:
            html = client.get_text(
                NEWS_SEARCH,
                params={
                    "SearchableText": "edital seleção residência",
                    "portal_type:list": "News Item",
                },
            )
            for item in self._parse_search_results(html):
                by_url[item["url"]] = item
        except Exception as exc:  # noqa: BLE001
            print(f"[{self.slug}] search: {exc}")

        return list(by_url.values())

    def _parse_cards(self, html: str, page_url: str) -> list[dict]:
        soup = BeautifulSoup(html, "lxml")
        items: list[dict] = []
        for a in soup.select("a.govbr-card-content[href], a.card[href], .govbr-cards a[href]"):
            href = a.get("href") or ""
            title_el = a.select_one(".titulo") or a
            title = title_el.get_text(" ", strip=True)
            if not href or not title:
                continue
            full = urljoin(page_url, href)
            if "/inca/" not in full:
                continue
            items.append({"title": unescape(title), "url": full, "summary": title})
        # Fallback: folder listing links under content-core
        if not items:
            core = soup.select_one("#content-core") or soup.select_one("#content")
            if core:
                for a in core.select("a[href]"):
                    href = a.get("href") or ""
                    title = a.get_text(" ", strip=True)
                    if not href or not title or len(title) < 8:
                        continue
                    if not OPEN_HINT_RE.search(title):
                        continue
                    full = urljoin(page_url, href)
                    if "/inca/" not in full:
                        continue
                    items.append({"title": unescape(title), "url": full, "summary": title})
        return items

    def _parse_content_links(self, html: str, page_url: str) -> list[dict]:
        soup = BeautifulSoup(html, "lxml")
        core = soup.select_one("#content-core") or soup.select_one("#main-content") or soup
        items: list[dict] = []
        for a in core.select("a[href]"):
            href = a.get("href") or ""
            title = a.get_text(" ", strip=True)
            if not href or not title or len(title) < 12:
                continue
            if not OPEN_HINT_RE.search(title):
                continue
            if SKIP_RE.search(title):
                continue
            full = urljoin(page_url, href)
            if "gov.br/inca" not in full:
                continue
            items.append({"title": unescape(title), "url": full, "summary": title})
        return items

    def _parse_search_results(self, html: str) -> list[dict]:
        soup = BeautifulSoup(html, "lxml")
        items: list[dict] = []
        for a in soup.select("a[href*='/assuntos/noticias/']"):
            href = a.get("href") or ""
            title = a.get_text(" ", strip=True)
            if not href or not title or len(title) < 12:
                continue
            if not OPEN_HINT_RE.search(title):
                continue
            items.append({"title": unescape(title), "url": href, "summary": title})
        return items

    def _fetch_detail(self, client: HttpClient, url: str) -> dict:
        html = client.get_text(url)
        soup = BeautifulSoup(html, "lxml")
        h1 = soup.select_one("h1") or soup.select_one("#content h1")
        page_title = h1.get_text(" ", strip=True) if h1 else None

        tab_sections = self._fetch_tab_sections(client, soup, url)
        core = soup.select_one("#content-core") or soup.select_one("#content")
        core_text = html_to_text(str(core)) if core else ""
        if LOADING_TAB_RE.search(core_text) and tab_sections:
            core_text = ""

        attachments = self._collect_attachments(soup, url)
        for _label, section_html, section_url in tab_sections:
            attachments.extend(
                self._collect_attachments(
                    BeautifulSoup(section_html, "lxml"), section_url
                )
            )
        # dedupe attachments preserving order
        seen_att: set[str] = set()
        uniq_att: list[tuple[str, str]] = []
        for label, href in attachments:
            if href in seen_att:
                continue
            seen_att.add(href)
            uniq_att.append((label, href))
        attachments = uniq_att

        blob_for_cpnu = " ".join(
            [page_title or "", core_text, url]
            + [t for t, _, _ in tab_sections]
            + [txt for _, html_s, _ in tab_sections for txt in [html_to_text(html_s)]]
        )
        cargos: list[dict] = []
        if CPNU_RE.search(blob_for_cpnu) or any(
            ATTACHMENT_EXT_RE.search(h) and "cargo" in (l + h).lower()
            for l, h in attachments
        ):
            cargos = self._load_cpnu_inca_cargos(client, attachments, blob_for_cpnu)

        description = self._compose_description(
            page_title=page_title,
            core_text=core_text,
            tab_sections=tab_sections,
            cargos=cargos,
            attachments=attachments,
        )
        requirements = self._compose_requirements(tab_sections, cargos)
        salary = salary_from_cargos(cargos)
        if not salary:
            m = MONEY_RE.search(description)
            salary = m.group(0) if m else None

        title = page_title
        if cargos:
            title = title_from_cargos(page_title, cargos)

        published = None
        time_el = soup.select_one("time[datetime], span.documentPublished, .documentPublished")
        if time_el:
            published = time_el.get("datetime") or time_el.get_text(" ", strip=True)

        specialty = "Oncologia" if (
            cargos or re.search(r"oncol|c[aâ]ncer|\binca\b|cpnu", blob_for_cpnu, re.I)
        ) else None

        return {
            "title": title,
            "description": description[:8000] if description else None,
            "requirements": requirements[:4000] if requirements else None,
            "salary": salary,
            "specialty": specialty,
            "published_at": published,
        }

    def _fetch_tab_sections(
        self, client: HttpClient, soup: BeautifulSoup, page_url: str
    ) -> list[tuple[str, str, str]]:
        """gov.br folders usam .tab-content[data-url] carregado por JS."""
        sections: list[tuple[str, str, str]] = []
        for pane in soup.select(".tab-content[data-url], [data-url]"):
            data_url = (pane.get("data-url") or "").strip()
            if not data_url:
                continue
            label = (pane.get("data-id") or "").strip() or data_url.rstrip("/").split("/")[-1]
            full = urljoin(page_url, data_url)
            try:
                html = client.get_text(full)
            except Exception as exc:  # noqa: BLE001
                print(f"[{self.slug}] aba {full}: {exc}")
                continue
            sections.append((label, html, full))
        return sections

    def _collect_attachments(
        self, soup: BeautifulSoup, page_url: str
    ) -> list[tuple[str, str]]:
        core = soup.select_one("#content-core") or soup.select_one("#content") or soup
        out: list[tuple[str, str]] = []
        for a in core.select("a[href]"):
            href = a.get("href") or ""
            if not href:
                continue
            full = _strip_view_suffix(urljoin(page_url, href))
            label = a.get_text(" ", strip=True) or full.rstrip("/").split("/")[-1]
            label = re.sub(r"\s*\(abre em outra aba\)\s*", "", label, flags=re.I).strip()
            if ATTACHMENT_EXT_RE.search(full):
                out.append((label, full))
            elif re.search(r"cargos?[_\s-]*sal[aá]rios?|\.xlsx?/view", href, re.I):
                out.append((label, full))
        return out

    def _load_cpnu_inca_cargos(
        self,
        client: HttpClient,
        attachments: list[tuple[str, str]],
        blob: str,
    ) -> list[dict]:
        xlsx_urls: list[str] = []
        for label, href in attachments:
            if re.search(r"\.xlsx?(?:$|[?#])", href, re.I) or "xlsx" in href.lower():
                xlsx_urls.append(_strip_view_suffix(href))
            elif re.search(r"cargo", f"{label} {href}", re.I) and href.endswith("/view"):
                xlsx_urls.append(_strip_view_suffix(href))

        # CPNU: ir à página MGI de cargos se ainda não há XLSX nos anexos INCA.
        if not xlsx_urls and CPNU_RE.search(blob):
            try:
                cargos_html = client.get_text(CPNU2_CARGOS_PAGE)
                for label, href in self._collect_attachments(
                    BeautifulSoup(cargos_html, "lxml"), CPNU2_CARGOS_PAGE
                ):
                    if re.search(r"\.xlsx?(?:$|[?#])", href, re.I) or "xlsx" in href:
                        xlsx_urls.append(_strip_view_suffix(href))
            except Exception as exc:  # noqa: BLE001
                print(f"[{self.slug}] cargos CPNU: {exc}")

        for xurl in xlsx_urls:
            try:
                data = client.get_bytes(xurl)
            except Exception as exc:  # noqa: BLE001
                print(f"[{self.slug}] xlsx {xurl}: {exc}")
                continue
            if not data or data[:2] != b"PK":
                # Pode ter vindo HTML da página /view
                continue
            cargos = parse_cpnu_cargos_xlsx(data)
            if cargos:
                return cargos
        return []

    def _compose_description(
        self,
        *,
        page_title: str | None,
        core_text: str,
        tab_sections: list[tuple[str, str, str]],
        cargos: list[dict],
        attachments: list[tuple[str, str]],
    ) -> str:
        parts: list[str] = []
        if page_title:
            parts.append(page_title)

        cargo_block = format_cargo_summary(cargos)
        if cargo_block:
            parts.append(cargo_block)

        # Preferir abas com conteúdo real; ignorar placeholders.
        for label, html, _url in tab_sections:
            text = html_to_text(
                str(
                    BeautifulSoup(html, "lxml").select_one("#content-core")
                    or BeautifulSoup(html, "lxml").select_one("#content")
                    or ""
                )
            )
            if not text or LOADING_TAB_RE.search(text):
                continue
            # Evitar repetir só o título da aba
            body = text.strip()
            if body.casefold() == label.casefold():
                continue
            parts.append(f"{label}\n{body}")

        if core_text and not LOADING_TAB_RE.search(core_text):
            parts.append(core_text)

        # Links úteis (PDF/XLS + páginas FGV/MGI/DOU referenciadas nas abas)
        outbound: list[str] = [f"- {label}: {href}" for label, href in attachments]
        for _label, html, section_url in tab_sections:
            soup = BeautifulSoup(html, "lxml")
            core = soup.select_one("#content-core") or soup
            for a in core.select("a[href]"):
                href = (a.get("href") or "").strip()
                text = a.get_text(" ", strip=True)
                text = re.sub(
                    r"\s*\(abre em outra aba\)\s*", "", text, flags=re.I
                ).strip()
                if not href or href.startswith("mailto:"):
                    continue
                full = href if href.startswith("http") else urljoin(section_url, href)
                low = full.lower()
                if any(
                    k in low
                    for k in (
                        "fgv.br",
                        "concursonacional",
                        "in.gov.br",
                        "inca.gov.br/publicacoes",
                        ".pdf",
                        ".xls",
                    )
                ):
                    outbound.append(f"- {text or full}: {full}")

        if outbound:
            seen: set[str] = set()
            lines = []
            for line in outbound:
                if line in seen:
                    continue
                seen.add(line)
                lines.append(line)
            parts.append("Documentos e links\n" + "\n".join(lines[:20]))

        return "\n\n".join(p for p in parts if p).strip()

    def _compose_requirements(
        self,
        tab_sections: list[tuple[str, str, str]],
        cargos: list[dict],
    ) -> str | None:
        chunks: list[str] = []
        if cargos:
            niveis = sorted(
                {
                    c["nivel"]
                    for c in cargos
                    if c.get("nivel")
                }
            )
            if niveis:
                chunks.append(
                    "Níveis exigidos no quadro INCA: " + "; ".join(niveis) + "."
                )
            chunks.append(
                "Requisitos e especialidades detalhados nos anexos do Edital ENAP "
                "nº 114/2025 (CPNU 2) e retificações — ver links na descrição."
            )
        for label, html, _url in tab_sections:
            if not re.search(r"document|admiss|ingresso|requisito", label, re.I):
                continue
            text = html_to_text(
                str(
                    BeautifulSoup(html, "lxml").select_one("#content-core")
                    or BeautifulSoup(html, "lxml").select_one("#content")
                    or ""
                )
            )
            if text and not LOADING_TAB_RE.search(text):
                chunks.append(text[:2500])
        if not chunks:
            return None
        return "\n\n".join(chunks).strip()
